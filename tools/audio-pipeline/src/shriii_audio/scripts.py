"""Per-language recording scripts for voice actors: XLSX (studio sheet) and printable HTML (VOX-07).

Each line carries: save-as file name, character, tone and direction notes, context, the en-IN reference
text, the script text in the target language (from an optional translations file exported by the ITC
workflow), a pronunciation field, timing targets (target, +-15 % window, 1.5x ceiling) and take boxes.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from shriii_audio.catalogue import Catalogue, Line, itc_rank
from shriii_audio.config import QcConfig
from shriii_audio.paths import default_direction_config

READY_ITC = "expert_reviewed"


@dataclass(frozen=True, slots=True)
class Direction:
    general: tuple[str, ...]
    tone: dict[str, str]
    speaker_role: dict[str, str]
    voice_forbidden_note: str


def load_direction(path: Path | None = None) -> Direction:
    raw = yaml.safe_load((path or default_direction_config()).read_text(encoding="utf-8"))
    return Direction(
        general=tuple(raw["general"]),
        tone=dict(raw["tone"]),
        speaker_role=dict(raw["speaker_role"]),
        voice_forbidden_note=str(raw["voice_forbidden_note"]).strip(),
    )


@dataclass(frozen=True, slots=True)
class Translation:
    text: str | None
    pronunciation: str | None


def load_translations(path: Path | None) -> dict[str, Translation]:
    """Optional ``{lines: {<line_id>: {text: ..., pronunciation: ...}}}`` file from the ITC workflow."""
    if path is None:
        return {}
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    lines = raw.get("lines", {})
    if not isinstance(lines, dict):
        raise ValueError(f"{path}: 'lines' must be a mapping of line id -> {{text, pronunciation}}")
    return {
        str(k): Translation(text=(v or {}).get("text"), pronunciation=(v or {}).get("pronunciation"))
        for k, v in lines.items()
    }


@dataclass(frozen=True, slots=True)
class ScriptRow:
    number: int
    line: Line
    save_as: str
    direction: str
    script_text: str
    pronunciation: str
    itc_state: str
    ready: bool
    target_s: float
    min_s: float
    max_s: float
    ceiling_s: float


def script_rows(
    catalogue: Catalogue,
    language: str,
    cfg: QcConfig,
    direction: Direction,
    translations: dict[str, Translation] | None = None,
    include_all: bool = False,
) -> list[ScriptRow]:
    if language not in catalogue.languages:
        raise ValueError(f"{language} is not a catalogue language ({', '.join(catalogue.languages)})")
    translations = translations or {}
    tol = cfg.duration.target_tolerance_fraction
    rows: list[ScriptRow] = []
    for line in catalogue.lines:
        rec = line.recording.get(language)
        if rec is not None and rec.status == "approved" and not include_all:
            continue
        itc = rec.itc_state if rec else "draft"
        tr = translations.get(line.id)
        if language == catalogue.reference_language:
            text = line.reference_text
            ready = True
        else:
            text = (tr.text if tr and tr.text else "") or ""
            ready = bool(text) and itc_rank(itc) >= itc_rank(READY_ITC)
        notes = [
            direction.speaker_role.get(line.speaker_role, line.speaker_role),
            direction.tone.get(line.tone, line.tone),
        ]
        if line.voice_forbidden:
            notes.append(direction.voice_forbidden_note)
        t = line.target_duration_ms / 1000
        rows.append(
            ScriptRow(
                number=len(rows) + 1,
                line=line,
                save_as=f"{line.id}.{language}.wav",
                direction=" ".join(notes),
                script_text=text,
                pronunciation=(tr.pronunciation if tr and tr.pronunciation else "") or "",
                itc_state=itc,
                ready=ready,
                target_s=round(t, 2),
                min_s=round(t * (1 - tol), 2),
                max_s=round(t * (1 + tol), 2),
                ceiling_s=round(t * cfg.duration.max_ratio_to_source, 2),
            )
        )
    return rows


COLUMNS: list[tuple[str, int]] = [
    ("#", 5),
    ("Line ID", 30),
    ("Save as", 38),
    ("Category", 11),
    ("Module", 14),
    ("Character", 11),
    ("Tone", 11),
    ("Direction", 48),
    ("Context", 40),
    ("Reference text (en-IN)", 48),
    ("Script text", 48),
    ("Pronunciation notes", 30),
    ("ITC state", 16),
    ("Ready to record", 10),
    ("Target s", 9),
    ("Window s (+-15%)", 14),
    ("Max s (1.5x)", 10),
    ("Never-voice", 10),
    ("Take 1", 9),
    ("Take 2", 9),
    ("Take 3", 9),
    ("Selected take", 10),
    ("Studio notes", 30),
]


def _cell(ws: Any, row: int, column: int, value: Any) -> Any:
    """Write a cell; text is always stored as text, so a script line starting with "=" never becomes a
    live formula in the actor's or reviewer's spreadsheet."""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def write_xlsx(rows: list[ScriptRow], language: str, catalogue: Catalogue, path: Path, direction: Direction) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = f"Script {language}"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3D5A80")
    wrap = Alignment(wrap_text=True, vertical="top")
    for idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name.replace("Script text", f"Script text ({language})"))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    take_dv = DataValidation(type="list", formula1='"ok,retake,noise,mistake,best"', allow_blank=True)
    ws.add_data_validation(take_dv)
    not_ready = PatternFill("solid", fgColor="FFF4E5")
    for r, row in enumerate(rows, start=2):
        values: list[Any] = [
            row.number,
            row.line.id,
            row.save_as,
            row.line.category,
            row.line.module_id or "",
            row.line.speaker_role,
            row.line.tone,
            row.direction,
            row.line.context,
            row.line.reference_text,
            row.script_text,
            row.pronunciation,
            row.itc_state,
            "yes" if row.ready else "no",
            row.target_s,
            f"{row.min_s}-{row.max_s}",
            row.ceiling_s,
            "YES" if row.line.voice_forbidden else "",
            "",
            "",
            "",
            "",
            "",
        ]
        for c, v in enumerate(values, start=1):
            cell = _cell(ws, r, c, v)
            cell.alignment = wrap
            if not row.ready:
                cell.fill = not_ready
        take_dv.add(f"S{r}:U{r}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    info = wb.create_sheet("Session")
    total_s = sum(r.target_s for r in rows)
    meta: list[tuple[str, Any]] = [
        ("Language", language),
        ("Catalogue version", catalogue.catalogue_version),
        ("Lines in script", len(rows)),
        ("Ready to record", sum(1 for r in rows if r.ready)),
        ("Total target audio (min)", round(total_s / 60, 1)),
        ("Estimated studio time (h, x12 rule)", round(total_s * 12 / 3600, 1)),
        ("Generated", datetime.now(UTC).replace(microsecond=0).isoformat()),
        ("", ""),
        ("General direction", ""),
        *[("", g) for g in direction.general],
    ]
    for i, (k, v) in enumerate(meta, start=1):
        _cell(info, i, 1, k).font = Font(bold=True)
        _cell(info, i, 2, v).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 36
    info.column_dimensions["B"].width = 100
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


HTML_CSS = """
:root{--bg:#fff;--fg:#1b1a17;--muted:#66625b;--line:#d9d4ca;--accent:#3d5a80;--warn:#a86400;--chip:#f3f0ea}
@media (prefers-color-scheme: dark){:root{--bg:#171614;--fg:#ece8e1;--muted:#a39d93;--line:#3a3631;
--accent:#8fb3e0;--warn:#f0a43a;--chip:#24221f}}
*{box-sizing:border-box}body{margin:0;padding:24px 16px;background:var(--bg);color:var(--fg);
font:15px/1.5 system-ui,-apple-system,"Segoe UI","Noto Sans","Noto Sans Devanagari","Noto Sans Tamil",sans-serif}
main{max-width:900px;margin:0 auto}h1{font-size:24px;margin:0}.muted{color:var(--muted)}
ul.general{padding-left:20px}
.card{border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin:14px 0;break-inside:avoid}
.card h3{margin:0 0 4px;font-size:15px;display:flex;justify-content:space-between;gap:12px;flex-wrap:wrap}
.card h3 code{font-size:14px}.chips{display:flex;flex-wrap:wrap;gap:6px;margin:6px 0}
.chip{background:var(--chip);border-radius:999px;padding:1px 10px;font-size:12px}
.chip.warn{color:var(--warn);font-weight:600}
.script{font-size:20px;line-height:1.5;margin:10px 0;padding:10px 12px;border-left:4px solid var(--accent);
background:var(--chip)}
.script.empty{font-size:14px;color:var(--warn)}
.ref{color:var(--muted)}dl{display:grid;grid-template-columns:150px 1fr;gap:4px 12px;margin:8px 0}
dt{color:var(--muted)}dd{margin:0}.field{border-bottom:1px dotted var(--muted);min-height:1.6em}
.takes{display:flex;gap:18px;margin-top:8px;flex-wrap:wrap}.takes span::before{content:"\\2610  "}
@media print{body{padding:0;background:#fff;color:#000}.card{page-break-inside:avoid}}
"""


def render_html(rows: list[ScriptRow], language: str, catalogue: Catalogue, direction: Direction) -> str:
    general = "".join(f"<li>{escape(g)}</li>" for g in direction.general)
    cards = []
    for r in rows:
        script = (
            f'<div class="script" lang="{escape(language)}">{escape(r.script_text)}</div>'
            if r.script_text
            else '<div class="script empty">Script text pending: awaiting the ITC-approved translation.</div>'
        )
        chips = [
            f'<span class="chip">{escape(r.line.category)}</span>',
            f'<span class="chip">{escape(r.line.speaker_role)}</span>',
            f'<span class="chip">tone: {escape(r.line.tone)}</span>',
            f'<span class="chip">ITC: {escape(r.itc_state)}</span>',
        ]
        if not r.ready:
            chips.append('<span class="chip warn">not ready to record</span>')
        if r.line.voice_forbidden:
            chips.append('<span class="chip warn">never-voice (headset only)</span>')
        pron = escape(r.pronunciation) if r.pronunciation else ""
        cards.append(
            f'<article class="card"><h3><span>{r.number}. <code>{escape(r.line.id)}</code></span>'
            f'<span class="muted">save as <code>{escape(r.save_as)}</code></span></h3>'
            f'<div class="chips">{"".join(chips)}</div>{script}'
            f'<div class="ref">Reference (en-IN): {escape(r.line.reference_text)}</div>'
            "<dl>"
            f"<dt>Direction</dt><dd>{escape(r.direction)}</dd>"
            f"<dt>Context</dt><dd>{escape(r.line.context)}</dd>"
            f"<dt>Timing</dt><dd>target {r.target_s} s &middot; window {r.min_s}&ndash;{r.max_s} s "
            f"&middot; never above {r.ceiling_s} s</dd>"
            f'<dt>Pronunciation</dt><dd class="field">{pron}</dd>'
            '<dt>Studio notes</dt><dd class="field"></dd>'
            "</dl>"
            '<div class="takes"><span>Take 1</span><span>Take 2</span><span>Take 3</span><span>Best: ____</span></div>'
            "</article>"
        )
    ready = sum(1 for r in rows if r.ready)
    return (
        f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>Recording Script {escape(language)}</title><style>{HTML_CSS}</style></head><body><main>"
        f"<h1>Recording script &middot; {escape(language)}</h1>"
        f'<p class="muted">Catalogue {escape(catalogue.catalogue_version)} &middot; {len(rows)} lines &middot; '
        f"{ready} ready to record &middot; total target {round(sum(r.target_s for r in rows) / 60, 1)} min</p>"
        f'<ul class="general">{general}</ul>{"".join(cards)}</main></body></html>'
    )


def write_scripts(
    catalogue: Catalogue,
    languages: list[str],
    cfg: QcConfig,
    out_dir: Path,
    *,
    direction: Direction | None = None,
    translations: dict[str, dict[str, Translation]] | None = None,
    include_all: bool = False,
) -> list[Path]:
    direction = direction or load_direction()
    written: list[Path] = []
    for lang in languages:
        rows = script_rows(catalogue, lang, cfg, direction, (translations or {}).get(lang), include_all)
        base = out_dir / "scripts"
        written.append(write_xlsx(rows, lang, catalogue, base / f"recording-script.{lang}.xlsx", direction))
        html_path = base / f"recording-script.{lang}.html"
        html_path.write_text(render_html(rows, lang, catalogue, direction), encoding="utf-8")
        written.append(html_path)
    return written
