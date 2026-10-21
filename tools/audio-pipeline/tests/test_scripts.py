"""Recording scripts from the real catalogue (read-only) and the catalogue loader."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from shriii_audio.catalogue import CatalogueError, load_catalogue
from shriii_audio.cli import main
from shriii_audio.config import QcConfig
from shriii_audio.contracts import ContractError
from shriii_audio.paths import default_catalogue
from shriii_audio.scripts import Translation, _cell, load_direction, script_rows, write_scripts


def test_real_catalogue_loads() -> None:
    cat = load_catalogue(default_catalogue())
    assert cat.reference_language == "en-IN"
    assert set(cat.languages) >= {"hi-IN", "mr-IN", "ta-IN"}
    assert len(cat.lines) > 100
    assert cat.get("attract.greeting") is not None


def test_invalid_catalogue_rejected(tmp_path: Path, catalogue_file: Path) -> None:
    data = yaml.safe_load(catalogue_file.read_text(encoding="utf-8"))
    data["lines"][0]["tone"] = "sarcastic"
    bad = tmp_path / "bad.yaml"
    bad.write_text(yaml.safe_dump(data), encoding="utf-8")
    with pytest.raises(ContractError):
        load_catalogue(bad)
    data["lines"][0]["tone"] = "warm"
    data["lines"].append(dict(data["lines"][0]))
    bad.write_text(yaml.safe_dump(data), encoding="utf-8")
    with pytest.raises(CatalogueError, match="duplicate"):
        load_catalogue(bad)


def test_direction_covers_every_enum() -> None:
    d = load_direction()
    cat = load_catalogue(default_catalogue())
    assert {line.tone for line in cat.lines} <= set(d.tone)
    assert {line.speaker_role for line in cat.lines} <= set(d.speaker_role)


def test_script_rows_reference_and_target_language(qc: QcConfig) -> None:
    cat = load_catalogue(default_catalogue())
    d = load_direction()
    en = script_rows(cat, "en-IN", qc, d)
    assert len(en) == len(cat.lines)  # nothing approved yet
    first = en[0]
    assert first.script_text == first.line.reference_text
    assert first.ready
    assert first.save_as == f"{first.line.id}.en-IN.wav"
    assert first.min_s == pytest.approx(first.target_s * 0.85, abs=0.01)
    assert first.ceiling_s == pytest.approx(first.target_s * 1.5, abs=0.01)
    hi = script_rows(cat, "hi-IN", qc, d)
    assert all(not r.ready and r.script_text == "" for r in hi)
    forbidden = [r for r in hi if r.line.voice_forbidden]
    assert forbidden and all("NEVER-VOICE" in r.direction for r in forbidden)
    # a translation alone is not enough: the ITC state must be expert_reviewed or later
    tr = {first.line.id: Translation(text="नमस्ते", pronunciation="na-MAS-tay")}
    hi2 = script_rows(cat, "hi-IN", qc, d, tr)
    assert hi2[0].script_text == "नमस्ते" and hi2[0].pronunciation == "na-MAS-tay"
    assert not hi2[0].ready
    with pytest.raises(ValueError, match="not a catalogue language"):
        script_rows(cat, "fr-FR", qc, d)


def test_script_ready_after_expert_review(qc: QcConfig, catalogue_file: Path) -> None:
    data = yaml.safe_load(catalogue_file.read_text(encoding="utf-8"))
    data["lines"][0]["recording"]["hi-IN"] = {"status": "scripted", "itc_state": "expert_reviewed"}
    data["lines"][1]["recording"]["hi-IN"] = {"status": "approved", "itc_state": "approved"}
    path = catalogue_file.parent / "lines2.yaml"
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    cat = load_catalogue(path)
    rows = script_rows(cat, "hi-IN", qc, load_direction(), {"guide.hello": Translation("नमस्ते", None)})
    assert rows[0].line.id == "guide.hello" and rows[0].ready
    assert "guide.bye" not in {r.line.id for r in rows}  # approved lines are skipped by default
    rows_all = script_rows(cat, "hi-IN", qc, load_direction(), include_all=True)
    assert "guide.bye" in {r.line.id for r in rows_all}


def test_write_scripts_xlsx_and_html(qc: QcConfig, tmp_path: Path) -> None:
    cat = load_catalogue(default_catalogue())
    paths = write_scripts(cat, ["hi-IN", "en-IN"], qc, tmp_path)
    assert {p.name for p in paths} == {
        "recording-script.hi-IN.xlsx",
        "recording-script.hi-IN.html",
        "recording-script.en-IN.xlsx",
        "recording-script.en-IN.html",
    }
    wb = load_workbook(tmp_path / "scripts" / "recording-script.hi-IN.xlsx")
    ws = wb["Script hi-IN"]
    headers = [c.value for c in ws[1]]
    for h in ("Line ID", "Save as", "Direction", "Pronunciation notes", "Target s", "Script text (hi-IN)", "Take 1"):
        assert h in headers
    assert ws.max_row == len(cat.lines) + 1
    save_as = headers.index("Save as")
    assert ws.cell(row=2, column=save_as + 1).value == f"{cat.lines[0].id}.hi-IN.wav"
    assert ws.freeze_panes == "C2"
    assert "Session" in wb.sheetnames
    html = (tmp_path / "scripts" / "recording-script.en-IN.html").read_text(encoding="utf-8")
    assert html.count('<article class="card">') == len(cat.lines)
    assert cat.lines[0].reference_text.split()[0] in html
    assert "@media print" in html


def test_cli_scripts(tmp_path: Path) -> None:
    tr = tmp_path / "hi.yaml"
    tr.write_text(
        yaml.safe_dump({"lines": {"attract.greeting": {"text": "नमस्ते!", "pronunciation": "nuh-MUS-tay"}}}),
        encoding="utf-8",
    )
    assert main(["scripts", "--out", str(tmp_path), "--language", "hi-IN", "--translations", f"hi-IN={tr}"]) == 0
    html = (tmp_path / "scripts" / "recording-script.hi-IN.html").read_text(encoding="utf-8")
    assert "नमस्ते!" in html and "nuh-MUS-tay" in html


def test_text_starting_with_equals_is_never_a_formula(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    _cell(ws, 1, 1, '=HYPERLINK("http://example.invalid","x")')
    _cell(ws, 1, 2, 1.5)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    back = load_workbook(path).active
    assert back is not None
    assert back.cell(1, 1).data_type == "s"
    assert back.cell(1, 1).value == '=HYPERLINK("http://example.invalid","x")'
    assert back.cell(1, 2).value == 1.5
