from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook  # type: ignore[import-untyped]
from shriii_i18n.cli import main
from shriii_i18n.repo import Repo
from shriii_i18n.spreadsheet import COLUMNS, SHEET

HEADER = {c.key: c.header for c in COLUMNS}


def _run(repo: Repo, *args: str) -> int:
    return main(["--root", str(repo.root), *args])


def test_validate_ok_and_failing(mini_repo: Repo, capsys: pytest.CaptureFixture[str]) -> None:
    assert _run(mini_repo, "validate", "--lang", "hi-IN") == 0
    assert "hi-IN: 0 error(s)" in capsys.readouterr().out
    path = mini_repo.translations_dir / "hi-IN.yaml"
    raw = yaml.safe_load(path.read_text(encoding="utf-8"))
    raw["entries"][0]["status"] = "approved"  # no history -> chain error, missing back-translation is present
    path.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    assert _run(mini_repo, "validate", "--lang", "hi-IN") == 1
    assert "history_status_mismatch" in capsys.readouterr().out


def test_unknown_language_is_an_error(mini_repo: Repo, capsys: pytest.CaptureFixture[str]) -> None:
    assert _run(mini_repo, "validate", "--lang", "xx-YY") == 1
    assert "unknown language" in capsys.readouterr().err


def test_export_import_roundtrip_via_cli(mini_repo: Repo, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "exports"
    assert _run(mini_repo, "export", "--lang", "hi-IN", "--out-dir", str(out)) == 0
    [wb_path] = list(out.glob("shriii-translations_hi-IN_*.xlsx"))
    wb = load_workbook(wb_path)
    ws = wb[SHEET]
    cols = {c.value: c.column for c in ws[1]}
    ws.cell(row=2, column=cols[HEADER["new_status"]], value="forward_translated")
    ws.cell(row=2, column=cols[HEADER["reviewer_name"]], value="Asha")
    ws.cell(row=2, column=cols[HEADER["reviewer_role"]], value="translator")
    ws.cell(row=2, column=cols[HEADER["review_date"]], value="2026-10-20")
    wb.save(wb_path)

    report = tmp_path / "report.json"
    assert _run(mini_repo, "import", str(wb_path), "--dry-run", "--report", str(report), "--today", "2026-10-21") == 0
    assert json.loads(report.read_text(encoding="utf-8"))[0]["transitions"] == {
        "draft_machine_needs_native_review -> forward_translated": 1
    }
    workflow = mini_repo.load_workflow()
    assert mini_repo.load_translations(workflow, "hi-IN").entries[0].status == "draft_machine_needs_native_review"

    assert _run(mini_repo, "import", str(wb_path), "--today", "2026-10-21") == 0
    assert mini_repo.load_translations(workflow, "hi-IN").entries[0].status == "forward_translated"
    assert _run(mini_repo, "validate", "--lang", "hi-IN") == 0

    # Re-importing the same (now stale) workbook is rejected as a concurrent change.
    assert _run(mini_repo, "import", str(wb_path), "--today", "2026-10-21") == 1
    assert "concurrent_change" in capsys.readouterr().out

    trail_dir = tmp_path / "trail"
    assert _run(mini_repo, "trail", "--lang", "hi-IN", "--out-dir", str(trail_dir)) == 0
    rows = list(csv.DictReader((trail_dir / "shriii-itc-trail_hi-IN.csv").open(encoding="utf-8-sig")))
    signed = [r for r in rows if r["by"]]
    assert signed == [
        {
            **signed[0],
            "line_id": "attract.touch_to_start",
            "to_status": "forward_translated",
            "role": "translator",
            "by": "Asha",
        }
    ]


def test_coverage_outputs_and_threshold(mini_repo: Repo, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "cov.json"
    assert _run(mini_repo, "coverage", "--lang", "hi-IN", "--format", "json", "--out", str(out)) == 0
    [cov] = json.loads(out.read_text(encoding="utf-8"))
    assert cov["entries"] == 3
    assert cov["milestones"]["draft_machine_needs_native_review"] == 1.0
    assert cov["milestones"]["forward_translated"] == 0.0
    assert cov["pack_ready"] == {"production": 0.0, "pilot": 0.0}
    capsys.readouterr()
    assert _run(mini_repo, "coverage", "--lang", "hi-IN") == 0
    md = capsys.readouterr().out
    assert md.startswith("# Localisation coverage report")
    assert "| Hindi (hi-IN) | 3/3 |" in md
    assert (
        _run(mini_repo, "coverage", "--lang", "hi-IN", "--min-status", "forward_translated", "--min-ratio", "0.5") == 1
    )


def test_sync_creates_missing_language_file(mini_repo: Repo, capsys: pytest.CaptureFixture[str]) -> None:
    assert not (mini_repo.translations_dir / "ta-IN.yaml").exists()
    assert _run(mini_repo, "sync", "--lang", "ta-IN") == 0
    data = mini_repo.load_translations(mini_repo.load_workflow(), "ta-IN")
    assert [e.id for e in data.entries] == [line.id for line in mini_repo.load_catalogue().lines]
    assert all(e.text == "" for e in data.entries)
    assert "+3 added" in capsys.readouterr().out
