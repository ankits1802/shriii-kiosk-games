"""Regression tests for verifier fixes: formula safety, stale-source rebase and coverage denominators."""

from __future__ import annotations

import csv
import datetime as dt
import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook  # type: ignore[import-untyped]
from shriii_i18n.checks import validate_translation_file
from shriii_i18n.cli import main
from shriii_i18n.models import HistoryEvent, TranslationFile
from shriii_i18n.repo import Repo
from shriii_i18n.spreadsheet import COLUMNS, SHEET, export_workbook, import_workbook
from shriii_i18n.sync import rebase_stale

TODAY = dt.date(2026, 10, 21)
HEADER = {c.key: c.header for c in COLUMNS}


def _run(repo: Repo, *args: str) -> int:
    return main(["--root", str(repo.root), *args])


def _set_entry(repo: Repo, line_id: str, **update: object) -> TranslationFile:
    workflow = repo.load_workflow()
    data = repo.load_translations(workflow, "hi-IN")
    entries = [e.model_copy(update=update) if e.id == line_id else e for e in data.entries]
    data = data.model_copy(update={"entries": entries})
    repo.save_translations(workflow, data)
    return data


def _export(repo: Repo, tmp_path: Path) -> Path:
    workflow = repo.load_workflow()
    return export_workbook(
        tmp_path / "hi.xlsx", repo.load_translations(workflow, "hi-IN"), repo.load_catalogue(), workflow
    )


def _change_source(repo: Repo, line_id: str, new_text: str) -> None:
    raw = yaml.safe_load(repo.catalogue_path.read_text(encoding="utf-8"))
    for line in raw["lines"]:
        if line["id"] == line_id:
            line["reference_text"] = new_text
    repo.catalogue_path.write_text(yaml.safe_dump(raw, allow_unicode=True, sort_keys=False), encoding="utf-8")


# --- spreadsheet formula safety ---------------------------------------------------------------------------------


def test_text_starting_with_equals_survives_roundtrip(mini_repo: Repo, tmp_path: Path) -> None:
    _set_entry(mini_repo, "feedback.well_done", translator_note="=SUM(1,2) is not a formula here")
    path = _export(mini_repo, tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    cols = {cell.value: cell.column for cell in ws[1]}
    note_cells = [ws.cell(row=r, column=cols[HEADER["translator_note"]]) for r in range(2, ws.max_row + 1)]
    assert all(c.data_type != "f" for c in note_cells)
    workflow = mini_repo.load_workflow()
    result = import_workbook(
        path, mini_repo.load_translations(workflow, "hi-IN"), mini_repo.load_catalogue(), workflow, today=TODAY
    )
    assert result.ok, [i.render() for i in result.issues.items]
    assert result.changes == []


def test_formula_typed_by_reviewer_is_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    cols = {cell.value: cell.column for cell in ws[1]}
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "feedback.well_done":
            ws.cell(row=r, column=cols[HEADER["text"]], value='=HYPERLINK("http://example.invalid","x")')
    wb.save(path)
    workflow = mini_repo.load_workflow()
    result = import_workbook(
        path, mini_repo.load_translations(workflow, "hi-IN"), mini_repo.load_catalogue(), workflow, today=TODAY
    )
    assert not result.ok
    assert {i.code for i in result.issues.errors} == {"formula_in_cell"}
    assert result.updated is None


def test_trail_csv_neutralises_formula_injection(mini_repo: Repo, tmp_path: Path) -> None:
    event = HistoryEvent(
        from_status="draft_machine_needs_native_review",
        to_status="flagged_for_revision",
        by="=cmd|' /C calc'!A0",
        role="translator",
        date=TODAY,
        comment="+fix this",
    )
    _set_entry(mini_repo, "feedback.well_done", status="flagged_for_revision", history=[event])
    assert _run(mini_repo, "trail", "--lang", "hi-IN", "--out-dir", str(tmp_path)) == 0
    rows = list(csv.DictReader((tmp_path / "shriii-itc-trail_hi-IN.csv").open(encoding="utf-8-sig")))
    [row] = [r for r in rows if r["line_id"] == "feedback.well_done"]
    assert row["by"] == "'=cmd|' /C calc'!A0"
    assert row["comment"] == "'+fix this"


# --- stale source rebase ----------------------------------------------------------------------------------------


def _signed_forward(repo: Repo, line_id: str) -> None:
    event = HistoryEvent(
        from_status="draft_machine_needs_native_review",
        to_status="forward_translated",
        by="Asha",
        role="translator",
        date=dt.date(2026, 10, 1),
    )
    _set_entry(repo, line_id, status="forward_translated", history=[event])


def test_rebase_stale_updates_drafts_and_flags_signed_entries(mini_repo: Repo) -> None:
    _signed_forward(mini_repo, "weighing_post.quantity")
    _change_source(mini_repo, "weighing_post.quantity", "The needle now shows {quantity}.")
    _change_source(mini_repo, "feedback.well_done", "Well done, friend!")
    workflow, catalogue = mini_repo.load_workflow(), mini_repo.load_catalogue()
    data = mini_repo.load_translations(workflow, "hi-IN")
    assert {i.code for i in validate_translation_file(data, catalogue, workflow).errors} == {"stale_source"}

    # Without a signer the signed-off entry cannot be flagged: nothing changes.
    unchanged, report = rebase_stale(data, catalogue, workflow, today=TODAY)
    assert not report.ok and unchanged == data

    updated, report = rebase_stale(data, catalogue, workflow, by="Meera", role="expert_reviewer", today=TODAY)
    assert report.ok
    assert report.rebased == ["feedback.well_done"]
    assert report.flagged == ["weighing_post.quantity"]
    entries = updated.by_id()
    assert entries["feedback.well_done"].status == workflow.initial_status
    flagged = entries["weighing_post.quantity"]
    assert flagged.status == "flagged_for_revision"
    assert flagged.history[-1].by == "Meera" and "now shows" in flagged.history[-1].comment
    assert validate_translation_file(updated, catalogue, workflow).ok


def test_sync_cli_rebase_stale(mini_repo: Repo, capsys: pytest.CaptureFixture[str]) -> None:
    _signed_forward(mini_repo, "weighing_post.quantity")
    _change_source(mini_repo, "weighing_post.quantity", "The needle now shows {quantity}.")
    assert _run(mini_repo, "validate", "--lang", "hi-IN") == 1
    assert _run(mini_repo, "sync", "--lang", "hi-IN", "--rebase-stale") == 1
    assert "rebase rejected" in capsys.readouterr().out
    assert (
        _run(
            mini_repo,
            "sync",
            "--lang",
            "hi-IN",
            "--rebase-stale",
            "--by",
            "Meera",
            "--role",
            "expert_reviewer",
            "--today",
            "2026-10-21",
        )
        == 0
    )
    assert _run(mini_repo, "validate", "--lang", "hi-IN") == 0
    data = mini_repo.load_translations(mini_repo.load_workflow(), "hi-IN")
    assert data.by_id()["weighing_post.quantity"].status == "flagged_for_revision"


# --- coverage denominators --------------------------------------------------------------------------------------


def test_coverage_milestones_exclude_unrecorded_lines(mini_repo: Repo, tmp_path: Path) -> None:
    _set_entry(mini_repo, "feedback.well_done", record=False, record_note="Not recorded in hi-IN.")
    for line_id in ("attract.touch_to_start", "weighing_post.quantity"):
        _signed_forward(mini_repo, line_id)
    out = tmp_path / "cov.json"
    assert _run(mini_repo, "coverage", "--lang", "hi-IN", "--format", "json", "--out", str(out)) == 0
    [cov] = json.loads(out.read_text(encoding="utf-8"))
    assert cov["to_record"] == 2
    assert cov["milestones"]["forward_translated"] == 1.0
    assert (
        _run(mini_repo, "coverage", "--lang", "hi-IN", "--min-status", "forward_translated", "--min-ratio", "1.0") == 0
    )
