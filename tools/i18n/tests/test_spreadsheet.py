from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook  # type: ignore[import-untyped]
from shriii_i18n.repo import Repo
from shriii_i18n.spreadsheet import (
    COLUMNS,
    LISTS_SHEET,
    META_SHEET,
    SHEET,
    ImportResult,
    WorkbookError,
    export_workbook,
    import_workbook,
)

TODAY = dt.date(2026, 10, 21)
HEADER = {c.key: c.header for c in COLUMNS}


def _export(repo: Repo, tmp_path: Path) -> Path:
    workflow = repo.load_workflow()
    return export_workbook(
        tmp_path / "out" / "hi.xlsx",
        repo.load_translations(workflow, "hi-IN"),
        repo.load_catalogue(),
        workflow,
        exported_at=dt.datetime(2026, 10, 21, 9, 0, tzinfo=dt.UTC),
    )


def _edit(path: Path, line_id: str, **values: object) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET]
    cols = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=cols[HEADER["line_id"]]).value == line_id:
            for key, value in values.items():
                ws.cell(row=row, column=cols[HEADER[key]], value=value)
            wb.save(path)
            return
    raise AssertionError(f"{line_id} not in workbook")


def _delete_row(path: Path, line_id: str) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == line_id:
            ws.delete_rows(row)
            wb.save(path)
            return
    raise AssertionError(line_id)


def _import(repo: Repo, path: Path, today: dt.date = TODAY) -> ImportResult:
    workflow = repo.load_workflow()
    guide = repo.load_register_guide()
    return import_workbook(
        path,
        repo.load_translations(workflow, "hi-IN"),
        repo.load_catalogue(),
        workflow,
        today=today,
        glossary=guide.glossary if guide else None,
    )


def _apply(repo: Repo, result: ImportResult) -> None:
    assert result.ok, [i.render() for i in result.issues.items]
    assert result.updated is not None
    repo.save_translations(repo.load_workflow(), result.updated)


def _codes(result: ImportResult) -> set[str]:
    return {i.code for i in result.issues.errors}


def _signoff(
    role: str, name: str = "A. Translator", date: object = "2026-10-20", comment: str = ""
) -> dict[str, object]:
    return {"reviewer_name": name, "reviewer_role": role, "review_date": date, "review_comment": comment}


# ---------------------------------------------------------------------------------------------------------------


def test_export_structure(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["README", SHEET, "status_guide", LISTS_SHEET, META_SHEET]
    assert wb[LISTS_SHEET].sheet_state == "hidden"
    assert wb[META_SHEET].sheet_state == "hidden"
    ws = wb[SHEET]
    assert [c.value for c in ws[1]] == [c.header for c in COLUMNS]
    assert ws.max_row == 4  # header + 3 catalogue lines
    assert ws.protection.sheet is True
    # editable columns unlocked, source columns locked
    cols = {cell.value: cell.column for cell in ws[1]}
    assert ws.cell(row=2, column=cols[HEADER["text"]]).protection.locked is False
    assert ws.cell(row=2, column=cols[HEADER["source_text"]]).protection.locked is True
    # dropdown validations reference the hidden lists sheet
    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert any(f and f.startswith(f"={LISTS_SHEET}!$A$2") for f in formulas)
    assert any(f and f.startswith(f"={LISTS_SHEET}!$B$2") for f in formulas)
    types = {dv.type for dv in ws.data_validations.dataValidation}
    assert {"list", "date"} <= types
    statuses = [c.value for c in wb[LISTS_SHEET]["A"][1:] if c.value]
    assert statuses[0] == "draft_machine_needs_native_review"
    assert "flagged_for_revision" in statuses
    meta = {r[0].value: r[1].value for r in wb[META_SHEET].iter_rows()}
    assert meta["language"] == "hi-IN"
    assert meta["catalogue_version"] == "9.9.9"


def test_unchanged_roundtrip_is_a_noop(mini_repo: Repo, tmp_path: Path) -> None:
    result = _import(mini_repo, _export(mini_repo, tmp_path))
    assert result.ok, [i.render() for i in result.issues.items]
    assert result.changes == []
    assert result.updated is not None
    assert result.updated == mini_repo.load_translations(mini_repo.load_workflow(), "hi-IN")


def test_forward_transition_with_signoff_is_applied(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(
        path,
        "attract.touch_to_start",
        text="शुरू करने के लिए, स्क्रीन को छुइए।",
        new_status="forward_translated",
        **_signoff("translator", date=dt.datetime(2026, 10, 20)),
    )
    result = _import(mini_repo, path)
    assert result.ok, [i.render() for i in result.issues.items]
    assert len(result.changes) == 1
    assert result.changes[0].changed_fields == ["status", "text"]
    assert result.updated is not None
    entry = result.updated.by_id()["attract.touch_to_start"]
    assert entry.status == "forward_translated"
    assert entry.text == "शुरू करने के लिए, स्क्रीन को छुइए।"
    assert entry.history[-1].role == "translator"
    assert entry.history[-1].date == dt.date(2026, 10, 20)
    assert entry.history[-1].source == "xlsx:hi.xlsx"


def test_missing_line_is_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _delete_row(path, "feedback.well_done")
    result = _import(mini_repo, path)
    assert "missing_line" in _codes(result)
    assert result.updated is None


def test_unknown_and_duplicate_lines_are_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    ws.cell(row=5, column=1, value="attract.touch_to_start")
    ws.cell(row=6, column=1, value="no.such_line")
    wb.save(path)
    codes = _codes(_import(mini_repo, path))
    assert {"duplicate_line", "unknown_line"} <= codes


def test_placeholder_mismatch_is_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(path, "weighing_post.quantity", text="सुई वज़न बता रही है।")
    result = _import(mini_repo, path)
    assert "placeholder_mismatch" in _codes(result)
    msg = next(i.message for i in result.issues.errors if i.code == "placeholder_mismatch")
    assert "{quantity}" in msg


def test_backward_transition_is_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(path, "feedback.well_done", new_status="forward_translated", **_signoff("translator"))
    assert _import(mini_repo, path).ok
    _apply(mini_repo, _import(mini_repo, path))
    path2 = _export(mini_repo, tmp_path / "second")
    _edit(path2, "feedback.well_done", new_status="draft_machine_needs_native_review", **_signoff("translator"))
    result = _import(mini_repo, path2)
    assert "illegal_transition" in _codes(result)
    assert "backward" in next(i.message for i in result.issues.errors if i.code == "illegal_transition")


def test_skipping_a_step_is_rejected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(path, "feedback.well_done", new_status="expert_reviewed", **_signoff("expert_reviewer"))
    result = _import(mini_repo, path)
    assert "illegal_transition" in _codes(result)


@pytest.mark.parametrize(
    ("signoff", "code"),
    [
        ({"reviewer_name": None, "reviewer_role": "translator", "review_date": "2026-10-20"}, "signoff_missing"),
        ({"reviewer_name": "X", "reviewer_role": "translator", "review_date": None}, "signoff_missing"),
        ({"reviewer_name": "X", "reviewer_role": "itc_lead", "review_date": "2026-10-20"}, "signoff_role"),
        ({"reviewer_name": "X", "reviewer_role": "translator", "review_date": "2026-12-01"}, "future_date"),
        ({"reviewer_name": "X", "reviewer_role": "translator", "review_date": "20/10/2026"}, "bad_date"),
    ],
)
def test_signoff_rules(mini_repo: Repo, tmp_path: Path, signoff: dict[str, object], code: str) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(path, "feedback.well_done", new_status="forward_translated", **signoff)
    assert code in _codes(_import(mini_repo, path))


def _advance(repo: Repo, tmp_path: Path, line_id: str, steps: list[tuple[str, str, str]]) -> None:
    for n, (status, role, name) in enumerate(steps):
        path = _export(repo, tmp_path / f"step{n}")
        _edit(repo_path := path, line_id, new_status=status, **_signoff(role, name=name))
        result = _import(repo, repo_path)
        assert result.ok, [i.render() for i in result.issues.items]
        _apply(repo, result)


def test_back_translator_must_be_independent(mini_repo: Repo, tmp_path: Path) -> None:
    _advance(mini_repo, tmp_path, "feedback.well_done", [("forward_translated", "translator", "Asha")])
    path = _export(mini_repo, tmp_path / "bt")
    _edit(
        path,
        "feedback.well_done",
        new_status="back_translated",
        back_translation="Very nice!",
        **_signoff("back_translator", name="asha "),
    )
    assert "signoff_not_independent" in _codes(_import(mini_repo, path))


def test_text_is_locked_after_back_translation(mini_repo: Repo, tmp_path: Path) -> None:
    _advance(
        mini_repo,
        tmp_path,
        "feedback.well_done",
        [("forward_translated", "translator", "Asha"), ("back_translated", "back_translator", "Ravi")],
    )
    path = _export(mini_repo, tmp_path / "edit")
    _edit(path, "feedback.well_done", text="बहुत अच्छा!")
    assert "locked_field" in _codes(_import(mini_repo, path))
    # ...but the reconciliation panel may edit on entering expert_reviewed
    path = _export(mini_repo, tmp_path / "panel")
    _edit(
        path,
        "feedback.well_done",
        text="बहुत अच्छा!",
        new_status="expert_reviewed",
        **_signoff("expert_reviewer", name="Panel"),
    )
    result = _import(mini_repo, path)
    assert result.ok, [i.render() for i in result.issues.items]


def test_flag_for_revision_needs_comment_and_restarts(mini_repo: Repo, tmp_path: Path) -> None:
    _advance(mini_repo, tmp_path, "feedback.well_done", [("forward_translated", "translator", "Asha")])
    path = _export(mini_repo, tmp_path / "flag")
    _edit(path, "feedback.well_done", new_status="flagged_for_revision", **_signoff("field_researcher", name="F"))
    assert "comment_missing" in _codes(_import(mini_repo, path))
    _edit(path, "feedback.well_done", review_comment="Sounds childish to elders")
    result = _import(mini_repo, path)
    assert result.ok
    _apply(mini_repo, result)
    path = _export(mini_repo, tmp_path / "restart")
    _edit(path, "feedback.well_done", new_status="back_translated", **_signoff("back_translator", name="R"))
    assert "illegal_transition" in _codes(_import(mini_repo, path))


def test_concurrent_yaml_change_is_detected(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    workflow = mini_repo.load_workflow()
    data = mini_repo.load_translations(workflow, "hi-IN")
    entries = [e.model_copy(update={"text": "बढ़िया!"}) if e.id == "feedback.well_done" else e for e in data.entries]
    mini_repo.save_translations(workflow, data.model_copy(update={"entries": entries}))
    assert "concurrent_change" in _codes(_import(mini_repo, path))


def test_stale_export_after_source_change(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    cat = yaml.safe_load(mini_repo.catalogue_path.read_text(encoding="utf-8"))
    cat["lines"][2]["reference_text"] = "Very well done!"
    mini_repo.catalogue_path.write_text(yaml.safe_dump(cat, allow_unicode=True), encoding="utf-8")
    assert "stale_export" in _codes(_import(mini_repo, path))


def test_signoff_without_transition_warns(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    _edit(path, "feedback.well_done", **_signoff("translator"))
    result = _import(mini_repo, path)
    assert result.ok
    assert "signoff_without_transition" in {i.code for i in result.issues.warnings}


def test_wrong_language_and_bad_workbook(mini_repo: Repo, tmp_path: Path) -> None:
    path = _export(mini_repo, tmp_path)
    wb = load_workbook(path)
    wb[META_SHEET]["B1"] = "something-else/1"
    wb.save(path)
    with pytest.raises(WorkbookError, match="schema"):
        _import(mini_repo, path)
    path = _export(mini_repo, tmp_path / "b")
    wb = load_workbook(path)
    wb[SHEET].cell(row=1, column=11, value="Renamed")
    wb.save(path)
    with pytest.raises(WorkbookError, match="missing columns"):
        _import(mini_repo, path)
