"""XLSX round-trip for translators and reviewers (VOX-07 "spreadsheet review").

Export writes one workbook per language: a protected ``translations`` sheet in which only the translator /
reviewer columns are editable, dropdown validation for the new status and the reviewer role, a date
validation for the review date, a ``status_guide`` sheet and hidden ``_lists`` / ``_meta`` sheets.

Import is all-or-nothing: the workbook is checked row by row against the catalogue, the current YAML and the
workflow (missing or unknown lines, stale exports, placeholder mismatches, forward-only transitions, field
locks, reviewer sign-off and signer independence) and the YAML is only rewritten when there are no errors.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import DataValidation  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from .checks import (
    IssueList,
    Signoff,
    check_entry_content,
    check_glossary,
    field_edit_allowed,
    transition_problems,
)
from .models import Catalogue, GlossaryTerm, HistoryEvent, TranslationEntry, TranslationFile, Workflow
from .text import estimate_duration, normalise, source_hash

WORKBOOK_SCHEMA = "shriii-i18n-xlsx/1"
SHEET = "translations"
META_SHEET = "_meta"
LISTS_SHEET = "_lists"
GUIDE_SHEET = "status_guide"
README_SHEET = "README"
TARGET_FONT = "Nirmala UI"  # ships with Windows and covers Devanagari and Tamil


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    width: int
    editable: bool = False
    hidden: bool = False
    wrap: bool = False
    target_script: bool = False


COLUMNS: tuple[Column, ...] = (
    Column("line_id", "Line ID", 34),
    Column("category", "Category", 12),
    Column("speaker_role", "Speaker", 10),
    Column("tone", "Tone", 12),
    Column("context", "Context (where this is heard)", 40, wrap=True),
    Column("target_duration_ms", "Target duration (ms)", 11),
    Column("voice_forbidden", "Headset only", 9),
    Column("record", "Record in this language", 10),
    Column("source_text", "Source text (en-IN)", 44, wrap=True),
    Column("current_status", "Current status", 22),
    Column("text", "Translation (recording script)", 52, editable=True, wrap=True, target_script=True),
    Column("back_translation", "Back-translation (English)", 44, editable=True, wrap=True),
    Column("register_note", "Register / tone note", 40, editable=True, wrap=True),
    Column("translator_note", "Translator note", 36, editable=True, wrap=True),
    Column("estimated_ms", "Estimated spoken ms (at export)", 12),
    Column("new_status", "New status", 26, editable=True),
    Column("reviewer_name", "Reviewer name", 20, editable=True),
    Column("reviewer_role", "Reviewer role", 18, editable=True),
    Column("review_date", "Review date (YYYY-MM-DD)", 14, editable=True),
    Column("review_comment", "Review comment", 36, editable=True, wrap=True),
    Column("source_hash", "source_hash", 18, hidden=True),
    Column("base_fingerprint", "base_fingerprint", 18, hidden=True),
)
COLUMN_INDEX = {c.key: i + 1 for i, c in enumerate(COLUMNS)}
HEADER_TO_KEY = {c.header: c.key for c in COLUMNS}

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_EDIT_FILL = PatternFill("solid", fgColor="FFF8E1")
_NORECORD_FILL = PatternFill("solid", fgColor="E7E6E6")


class WorkbookError(ValueError):
    """The workbook is not a usable shriii-i18n export (wrong schema, missing sheets or columns)."""


def entry_fingerprint(entry: TranslationEntry) -> str:
    """Fingerprint of the editable state of an entry at export time (detects concurrent YAML changes)."""
    payload = json.dumps(
        [entry.status, normalise(entry.text), normalise(entry.back_translation), len(entry.history)],
        ensure_ascii=False,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


# --------------------------------------------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------------------------------------------


def export_workbook(
    path: Path,
    data: TranslationFile,
    catalogue: Catalogue,
    workflow: Workflow,
    *,
    exported_at: dt.datetime | None = None,
) -> Path:
    lang = workflow.language(data.language)
    exported_at = exported_at or dt.datetime.now(dt.UTC).replace(microsecond=0)
    entries = data.by_id()
    wb = Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = README_SHEET
    ws = wb.create_sheet(SHEET)
    guide = wb.create_sheet(GUIDE_SHEET)
    lists = wb.create_sheet(LISTS_SHEET)
    meta = wb.create_sheet(META_SHEET)

    _write_readme(readme, lang.tag, lang.name, workflow)
    _write_guide(guide, workflow)
    n_status, n_roles = _write_lists(lists, workflow)
    _write_meta(meta, data, catalogue, exported_at)

    for col in COLUMNS:
        cell = ws.cell(row=1, column=COLUMN_INDEX[col.key], value=col.header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = get_column_letter(COLUMN_INDEX[col.key])
        ws.column_dimensions[letter].width = col.width
        ws.column_dimensions[letter].hidden = col.hidden
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "B2"

    for r, line in enumerate(catalogue.lines, start=2):
        entry = entries.get(line.id)
        if entry is None:
            raise WorkbookError(f"{data.language}: no entry for {line.id}; run `shriii-i18n sync` before exporting")
        est = (
            estimate_duration(entry.text, lang.script, lang.tag, line.target_duration_ms, workflow.quality).estimated_ms
            if entry.text
            else None
        )
        values: dict[str, Any] = {
            "line_id": line.id,
            "category": line.category,
            "speaker_role": line.speaker_role,
            "tone": line.tone,
            "context": line.context,
            "target_duration_ms": line.target_duration_ms,
            "voice_forbidden": "yes" if line.voice_forbidden else "no",
            "record": "yes" if entry.record else f"no: {entry.record_note}",
            "source_text": line.reference_text,
            "current_status": entry.status,
            "text": entry.text,
            "back_translation": entry.back_translation,
            "register_note": entry.register_note,
            "translator_note": entry.translator_note,
            "estimated_ms": est,
            "new_status": entry.status,
            "reviewer_name": None,
            "reviewer_role": None,
            "review_date": None,
            "review_comment": None,
            "source_hash": entry.source_hash,
            "base_fingerprint": entry_fingerprint(entry),
        }
        for col in COLUMNS:
            cell = ws.cell(row=r, column=COLUMN_INDEX[col.key], value=values[col.key])
            _force_text(cell)
            cell.alignment = Alignment(wrap_text=col.wrap, vertical="top")
            if col.editable:
                cell.protection = Protection(locked=False)
                cell.fill = _EDIT_FILL
            elif not entry.record:
                cell.fill = _NORECORD_FILL
            if col.target_script:
                cell.font = Font(name=TARGET_FONT, size=12)
            if col.key == "review_date":
                cell.number_format = "yyyy-mm-dd"

    last = len(catalogue.lines) + 1
    status_dv = DataValidation(
        type="list",
        formula1=f"={LISTS_SHEET}!$A$2:$A${n_status + 1}",
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Unknown status",
        error="Pick a status from the list (see the status_guide sheet).",
        promptTitle="New status",
        prompt="Leave unchanged, or move ONE step forward (or to flagged_for_revision). Sign off on the right.",
        showInputMessage=True,
    )
    role_dv = DataValidation(
        type="list",
        formula1=f"={LISTS_SHEET}!$B$2:$B${n_roles + 1}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Unknown role",
        error="Pick your role from the list.",
    )
    date_dv = DataValidation(
        type="date",
        operator="greaterThan",
        formula1="DATE(2024,1,1)",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Review date",
        error="Enter a date (YYYY-MM-DD).",
    )
    for dv, key in ((status_dv, "new_status"), (role_dv, "reviewer_role"), (date_dv, "review_date")):
        letter = get_column_letter(COLUMN_INDEX[key])
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    lists.sheet_state = "hidden"
    meta.sheet_state = "hidden"
    wb.active = wb.index(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _force_text(cell: Any) -> None:
    """Store strings literally: openpyxl turns any string starting with '=' into a formula, which would be
    lost on re-import (no cached value) and could execute in the reviewer's spreadsheet program."""
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


def _write_readme(ws: Worksheet, tag: str, name: str, workflow: Workflow) -> None:
    lines = [
        f"SHRIII recording-script review workbook - {name} ({tag})",
        "",
        "What to do",
        "1. Work only in the yellow columns of the 'translations' sheet. Grey rows are not recorded in this language.",
        "2. Write the Translation as it should be SPOKEN by a native voice artist: warm, respectful, simple, no "
        "words that need schooling or reading to understand. Do not add or remove meaning.",
        "3. Back-translation: a literal English rendering of the Translation, written by someone who did not "
        "see the English source (independent back-translation).",
        "4. To move a line forward, pick the next status in 'New status' and fill Reviewer name, Reviewer role, "
        "Review date. One step per round trip. 'flagged_for_revision' needs a Review comment.",
        "5. Do not edit hidden columns, rename sheets, or re-order the header row. Send the file back as .xlsx.",
        "",
        "Status ladder (see 'status_guide' for who signs each step):",
        " -> ".join(s.id for s in workflow.forward_path()),
        "Any step can go to flagged_for_revision; a flagged line restarts at forward_translated.",
        "",
        "The import rejects the whole workbook if any row breaks a rule; the error report names the row.",
    ]
    for r, text in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1:
            cell.font = Font(bold=True, size=14)
        elif text in ("What to do",):
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 120


def _write_guide(ws: Worksheet, workflow: Workflow) -> None:
    headers = ["Status", "Label", "Who signs the move into it", "Requires", "Allowed next", "Meaning"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    for r, s in enumerate(workflow.statuses, start=2):
        signer = "any listed role" if s.signoff_role == "*" else (s.signoff_role or "-")
        row = [
            s.id,
            s.label,
            signer,
            ", ".join(s.requires) or "-",
            ", ".join(workflow.transitions.get(s.id, [])),
            s.description,
        ]
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate([30, 26, 22, 26, 40, 70], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _write_lists(ws: Worksheet, workflow: Workflow) -> tuple[int, int]:
    ws.cell(row=1, column=1, value="status")
    ws.cell(row=1, column=2, value="role")
    for r, s in enumerate(workflow.statuses, start=2):
        ws.cell(row=r, column=1, value=s.id)
    for r, role in enumerate(workflow.roles, start=2):
        ws.cell(row=r, column=2, value=role)
    return len(workflow.statuses), len(workflow.roles)


def _write_meta(ws: Worksheet, data: TranslationFile, catalogue: Catalogue, exported_at: dt.datetime) -> None:
    rows = [
        ("workbook_schema", WORKBOOK_SCHEMA),
        ("language", data.language),
        ("catalogue_version", catalogue.catalogue_version),
        ("exported_at", exported_at.isoformat()),
        ("columns", ",".join(c.key for c in COLUMNS)),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)


# --------------------------------------------------------------------------------------------------------------
# Import
# --------------------------------------------------------------------------------------------------------------


@dataclass
class RowChange:
    line_id: str
    row: int
    from_status: str
    to_status: str
    changed_fields: list[str]


@dataclass
class ImportResult:
    language: str
    workbook: Path
    issues: IssueList
    changes: list[RowChange] = field(default_factory=list)
    updated: TranslationFile | None = None

    @property
    def ok(self) -> bool:
        return self.issues.ok

    def summary(self) -> dict[str, object]:
        transitions: dict[str, int] = {}
        for ch in self.changes:
            if ch.from_status != ch.to_status:
                key = f"{ch.from_status} -> {ch.to_status}"
                transitions[key] = transitions.get(key, 0) + 1
        return {
            "language": self.language,
            "workbook": str(self.workbook),
            "ok": self.ok,
            "errors": len(self.issues.errors),
            "warnings": len(self.issues.warnings),
            "rows_changed": len(self.changes),
            "transitions": transitions,
            "issues": [i.as_dict() for i in self.issues.items],
        }


def read_workbook_language(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        meta = _read_meta(wb)
    finally:
        wb.close()
    return meta["language"]


def _read_meta(wb: Any) -> dict[str, str]:
    if META_SHEET not in wb.sheetnames:
        raise WorkbookError(f"workbook has no {META_SHEET!r} sheet; is it a shriii-i18n export?")
    meta: dict[str, str] = {}
    for row in wb[META_SHEET].iter_rows(values_only=True):
        if row and row[0] is not None:
            meta[str(row[0])] = "" if len(row) < 2 or row[1] is None else str(row[1])
    if meta.get("workbook_schema") != WORKBOOK_SCHEMA:
        raise WorkbookError(f"unsupported workbook schema {meta.get('workbook_schema')!r} (expected {WORKBOOK_SCHEMA})")
    if not meta.get("language"):
        raise WorkbookError("workbook meta has no language")
    return meta


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_date(value: object) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    try:
        return dt.date.fromisoformat(text[:10])
    except ValueError as exc:
        raise ValueError(f"review date {text!r} is not YYYY-MM-DD") from exc


def import_workbook(
    path: Path,
    data: TranslationFile,
    catalogue: Catalogue,
    workflow: Workflow,
    *,
    today: dt.date | None = None,
    glossary: list[GlossaryTerm] | None = None,
) -> ImportResult:
    today = today or dt.date.today()
    issues = IssueList()
    result = ImportResult(language=data.language, workbook=path, issues=issues)
    # data_only=False so that formulas typed by a reviewer are seen (and rejected) instead of silently read
    # back as their cached value or as empty.
    wb = load_workbook(path, data_only=False)
    meta = _read_meta(wb)
    lang = workflow.language(data.language)
    where_lang = {"language": lang.tag}
    if meta["language"] != data.language:
        raise WorkbookError(f"workbook is for {meta['language']}, not {data.language}")
    if meta.get("catalogue_version") != catalogue.catalogue_version:
        issues.warning(
            "catalogue_version",
            f"workbook exported from catalogue {meta.get('catalogue_version')}, "
            f"current is {catalogue.catalogue_version}",
            **where_lang,
        )
    if SHEET not in wb.sheetnames:
        raise WorkbookError(f"workbook has no {SHEET!r} sheet")
    ws = wb[SHEET]
    header = [_cell_text(c.value) for c in ws[1]]
    col_of: dict[str, int] = {}
    for idx, label in enumerate(header):
        key = HEADER_TO_KEY.get(label)
        if key:
            col_of[key] = idx
    missing_cols = [c.header for c in COLUMNS if c.key not in col_of]
    if missing_cols:
        raise WorkbookError(f"workbook is missing columns: {', '.join(missing_cols)}")

    lines = catalogue.by_id()
    entries = data.by_id()
    seen: dict[str, int] = {}
    staged: dict[str, TranslationEntry] = {}

    for r, cells in enumerate(ws.iter_rows(min_row=2), start=2):
        row = {k: (cells[i].value if i < len(cells) else None) for k, i in col_of.items()}
        if all(_cell_text(v) == "" for v in row.values()):
            continue
        line_id = _cell_text(row["line_id"])
        where: dict[str, object] = {"language": lang.tag, "line_id": line_id, "row": r}
        if not line_id:
            issues.error("missing_line_id", "row has content but no Line ID", language=lang.tag, row=r)
            continue
        if line_id in seen:
            issues.error("duplicate_line", f"line already appears in row {seen[line_id]}", **where)
            continue
        seen[line_id] = r
        formula_cols = [
            COLUMNS[COLUMN_INDEX[k] - 1].header
            for k, i in col_of.items()
            if i < len(cells) and cells[i].data_type == "f"
        ]
        if formula_cols:
            issues.error(
                "formula_in_cell",
                f"cells hold spreadsheet formulas ({', '.join(formula_cols)}); type plain text instead",
                **where,
            )
            continue
        line = lines.get(line_id)
        entry = entries.get(line_id)
        if line is None or entry is None:
            issues.error("unknown_line", "Line ID is not in the catalogue / translation file", **where)
            continue
        if _cell_text(row["source_hash"]) != source_hash(line.reference_text):
            issues.error("stale_export", "the source text changed after this workbook was exported; re-export", **where)
            continue
        if _cell_text(row["base_fingerprint"]) != entry_fingerprint(entry):
            issues.error(
                "concurrent_change",
                "this line changed in git after the workbook was exported; re-export and redo this row",
                **where,
            )
            continue

        new_status = _cell_text(row["new_status"]) or entry.status
        new_text = normalise(_cell_text(row["text"]))
        new_bt = normalise(_cell_text(row["back_translation"]))
        new_register = _cell_text(row["register_note"])
        new_tnote = _cell_text(row["translator_note"])
        reviewer = _cell_text(row["reviewer_name"])
        role = _cell_text(row["reviewer_role"])
        comment = _cell_text(row["review_comment"])
        try:
            review_date = _cell_date(row["review_date"])
        except ValueError as exc:
            issues.error("bad_date", str(exc), **where)
            continue

        if new_status not in workflow.status_ids():
            issues.error("unknown_status", f"new status {new_status!r} is not declared", **where)
            continue

        changed: list[str] = []
        history = list(entry.history)
        if new_status != entry.status:
            prior = {ev.to_status: ev.by for ev in entry.history}
            for code, msg in transition_problems(
                workflow, entry.status, new_status, Signoff(reviewer, role, comment), prior
            ):
                issues.error(code, msg, **where)
            if review_date is None:
                issues.error("signoff_missing", f"moving to {new_status} needs a review date", **where)
            else:
                if review_date > today:
                    issues.error("future_date", f"review date {review_date} is in the future", **where)
                if entry.history and review_date < entry.history[-1].date:
                    issues.error("history_date_order", "review date is before the previous sign-off", **where)
                history.append(
                    HistoryEvent(
                        from_status=entry.status,
                        to_status=new_status,
                        by=reviewer or "?",
                        role=role or "?",
                        date=review_date,
                        comment=comment,
                        source=f"xlsx:{path.name}",
                    )
                )
            changed.append("status")
        elif reviewer or role or review_date or comment:
            issues.warning(
                "signoff_without_transition",
                "reviewer fields are filled but the status did not change; the sign-off was ignored",
                **where,
            )

        for field_name, new_value, old_value in (
            ("text", new_text, normalise(entry.text)),
            ("back_translation", new_bt, normalise(entry.back_translation)),
        ):
            if new_value != old_value:
                if not field_edit_allowed(workflow, field_name, entry.status, new_status):
                    issues.error(
                        "locked_field",
                        f"{field_name} cannot change on {entry.status} -> {new_status}; "
                        "flag the line for revision first",
                        **where,
                    )
                changed.append(field_name)
        if new_register != entry.register_note:
            changed.append("register_note")
        if new_tnote != entry.translator_note:
            changed.append("translator_note")
        if (
            "text" in changed
            and "back_translation" not in changed
            and new_bt
            and new_status != workflow.initial_status
            and field_edit_allowed(workflow, "back_translation", entry.status, new_status)
        ):
            issues.warning(
                "back_translation_not_updated",
                "translation changed but the back-translation did not; check it still matches",
                **where,
            )

        candidate = entry.model_copy(
            update={
                "status": new_status,
                "text": new_text,
                "back_translation": new_bt,
                "register_note": new_register,
                "translator_note": new_tnote,
                "history": history,
            }
        )
        if changed:
            # Unchanged rows were validated when they entered git; only re-check what this workbook touches.
            issues.extend(check_entry_content(candidate, line, lang, workflow, row=r))
            if glossary:
                issues.extend(check_glossary(candidate, line, lang.tag, glossary, row=r))
            staged[line_id] = candidate
            result.changes.append(RowChange(line_id, r, entry.status, new_status, changed))

    for line in catalogue.lines:
        if line.id not in seen:
            issues.error(
                "missing_line", "catalogue line is missing from the workbook", language=lang.tag, line_id=line.id
            )

    if issues.ok:
        result.updated = data.model_copy(update={"entries": [staged.get(e.id, e) for e in data.entries]})
    wb.close()
    return result
