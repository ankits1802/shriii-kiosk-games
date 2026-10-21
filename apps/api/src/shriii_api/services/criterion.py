"""Criterion battery (layer 3): instrument registry, response capture, scoring, imports (CRT-02..05).

* The registry is loaded from ``content/instruments/*.yaml`` (structure, keys, licensing metadata;
  never copyrighted wording).
* Administration requires: an approved language version (409 otherwise, CRT-02), licensing
  permission granted or not required (CRT-03), criterion consent for the pid, and never a public
  kiosk context (CRT-04: read-aloud records from ``open_public`` are rejected).
* Responses are append-only; corrections are superseding rows with a reason; every (re)scoring creates
  a new scoring run and marks the previous scale scores non-current (audit history).
"""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from shriii_api import content
from shriii_api.adapters.scoring import ResponseInput, score_instrument, split_measure
from shriii_api.canonical import canonical_sha256, sha256_hex
from shriii_api.db.models.criterion import (
    ImportBatch,
    ImportRow,
    Instrument,
    InstrumentVersion,
    Item,
    LanguageVersion,
    Permission,
    Response,
    ScaleScore,
    ScoringKey,
)
from shriii_api.db.models.participants import Participant
from shriii_api.db.models.scored import ScoringRun
from shriii_api.errors import DomainError
from shriii_api.security import is_valid_code, normalise_code
from shriii_api.services import consent as consent_svc

PERMITTED_STATUSES = frozenset({"granted", "not_required"})


class CriterionError(DomainError):
    """Criterion error."""


# ---------------------------------------------------------------- registry


def sync_instruments(db: Session, actor: str) -> dict[str, int]:
    """Load/refresh the registry from YAML. New versions are added; existing versions are immutable."""
    stats = {"instruments": 0, "versions_added": 0, "unchanged": 0}
    for doc in content.instrument_documents():
        d = doc.data
        inst = db.get(Instrument, d["instrument_id"])
        if inst is None:
            inst = Instrument(
                instrument_id=d["instrument_id"],
                name=d["name"],
                short_name=d["short_name"],
                kind=d["kind"],
                citation=d["citation"],
                constructs_measured=d["constructs_measured"],
                current_version=d["version"],
            )
            db.add(inst)
            db.flush()
        stats["instruments"] += 1
        existing = db.execute(
            select(InstrumentVersion).where(
                InstrumentVersion.instrument_id == d["instrument_id"], InstrumentVersion.version == d["version"]
            )
        ).scalar_one_or_none()
        if existing is not None:
            stats["unchanged"] += 1
            continue
        inst.name, inst.short_name, inst.kind = d["name"], d["short_name"], d["kind"]
        inst.citation, inst.constructs_measured, inst.current_version = (
            d["citation"],
            d["constructs_measured"],
            d["version"],
        )
        version = InstrumentVersion(
            instrument_id=d["instrument_id"],
            version=d["version"],
            key_version=d["scoring"]["key_version"],
            definition=d,
            sha256=doc.sha256,
            status="active",
            source_path=doc.relpath,
        )
        db.add(version)
        db.flush()
        for it in d["items"]:
            db.add(
                Item(
                    instrument_version_id=version.id,
                    item_id=it["item_id"],
                    number=it["number"],
                    subscales=it["subscales"],
                    reverse_keyed=it["reverse_keyed"],
                    attention_check=it["attention_check"],
                    block=it.get("block"),
                    item_text_source=it["item_text_source"],
                )
            )
        for lv in d["language_versions"]:
            db.add(
                LanguageVersion(
                    instrument_version_id=version.id,
                    language=lv["language"],
                    itc_state=lv["itc_state"],
                    approved=bool(lv["approved"]),
                    translation_permission=lv["translation_permission"],
                )
            )
        key_doc = {
            "scales": d["subscales"],
            "scoring": d["scoring"],
            "reverse_formula": d["response_scale"]["reverse_formula"],
            "reverse_items": [i["item_id"] for i in d["items"] if i["reverse_keyed"]],
        }
        db.add(
            ScoringKey(
                instrument_version_id=version.id,
                key_version=d["scoring"]["key_version"],
                prorate_min_prop=float(d["scoring"]["prorate_min_prop"]),
                reverse_formula=d["response_scale"]["reverse_formula"],
                reverse_items=key_doc["reverse_items"],
                scales=d["subscales"],
                sha256=canonical_sha256(key_doc),
            )
        )
        lic = d["licensing"]
        db.add(
            Permission(
                instrument_id=d["instrument_id"],
                language=None,
                holder=lic["holder"],
                licence=lic["licence"],
                permission_status=lic["permission_status"],
                permission_document=lic["permission_document"],
                allowed_use=lic["allowed_use"],
                expiry=date.fromisoformat(lic["expiry"]) if lic.get("expiry") else None,
                commercial_use=bool(lic["commercial_use"]),
                notes=lic.get("notes"),
                recorded_by=actor,
            )
        )
        stats["versions_added"] += 1
    db.flush()
    return stats


def get_version(db: Session, instrument_id: str, version: str | None = None) -> InstrumentVersion:
    stmt = select(InstrumentVersion).where(InstrumentVersion.instrument_id == instrument_id)
    if version is not None:
        stmt = stmt.where(InstrumentVersion.version == version)
    else:
        inst = db.get(Instrument, instrument_id)
        if inst is None:
            raise CriterionError(404, "unknown_instrument", f"instrument {instrument_id} not found")
        stmt = stmt.where(InstrumentVersion.version == inst.current_version)
    row = db.execute(stmt).scalar_one_or_none()
    if row is None:
        raise CriterionError(404, "unknown_instrument_version", f"{instrument_id}@{version} not found")
    return row


def current_permission(db: Session, instrument_id: str) -> Permission | None:
    return db.execute(
        select(Permission)
        .where(Permission.instrument_id == instrument_id, Permission.language.is_(None))
        .order_by(Permission.created_at.desc(), Permission.id.desc())
        .limit(1)
    ).scalar_one_or_none()


def permission_warnings(perm: Permission | None, today: date | None = None) -> list[str]:
    today = today or datetime.now(UTC).date()
    warnings: list[str] = []
    if perm is None:
        return ["no_permission_record"]
    if perm.permission_status not in PERMITTED_STATUSES:
        warnings.append(f"permission_{perm.permission_status}")
    if perm.expiry is not None:
        days = (perm.expiry - today).days
        if days < 0:
            warnings.append("permission_expired")
        elif days <= 60:
            warnings.append("permission_expires_within_60_days")
    return warnings


def assert_administrable(db: Session, version: InstrumentVersion, language: str) -> LanguageVersion:
    lv = db.execute(
        select(LanguageVersion).where(
            LanguageVersion.instrument_version_id == version.id, LanguageVersion.language == language
        )
    ).scalar_one_or_none()
    if lv is None:
        raise CriterionError(409, "language_version_missing", f"no {language} version of {version.instrument_id}")
    if not lv.approved:
        raise CriterionError(
            409,
            "language_version_not_approved",
            f"{version.instrument_id}@{version.version} ({language}) is not approved",
        )
    perm = current_permission(db, version.instrument_id)
    if (
        perm is None
        or perm.permission_status not in PERMITTED_STATUSES
        or (perm.expiry is not None and perm.expiry < datetime.now(UTC).date())
    ):
        raise CriterionError(409, "licence_not_granted", f"{version.instrument_id} has no granted permission")
    return lv


# ---------------------------------------------------------------- responses


@dataclass(slots=True)
class ItemResponse:
    item_id: str
    raw_response: float | None
    response_latency_ms: float | None = None
    expected_response: float | None = None


@dataclass(slots=True)
class Administration:
    pid: str
    study_id: str
    wave: int
    occasion_id: str
    instrument_id: str
    instrument_version: str | None
    language: str
    block: str
    administration_mode: str
    administrator_id: str | None
    device_id: str
    audio_version: str | None
    site_mode: str
    rater_role: str | None
    rated_pid: str | None
    source: str
    planned_missing_items: list[str]
    responses: list[ItemResponse]
    import_id: uuid.UUID | None = None


def _resolve_item(definition: dict[str, Any], item_id: str) -> tuple[dict[str, Any], str | None]:
    """The registry item for ``item_id`` plus the performance measure (``<item>.<measure>``), if any."""
    items = {i["item_id"]: i for i in definition["items"]}
    base, measure = split_measure(items, item_id)
    item = items.get(base)
    if item is None:
        raise CriterionError(422, "unknown_item", f"{item_id} is not part of {definition['instrument_id']}")
    if definition.get("kind") == "performance":
        if measure is None or measure not in (item.get("measures") or []):
            raise CriterionError(
                422,
                "performance_measure_required",
                f"performance responses use <item>.<measure>; {base} measures: {item.get('measures') or []}",
            )
    elif measure is not None:
        raise CriterionError(422, "unknown_item", f"{item_id} is not part of {definition['instrument_id']}")
    return item, measure


def _validate_value(instrument: dict[str, Any], item_id: str, value: float | None) -> None:
    if value is None:
        return
    if instrument.get("kind") == "performance":
        if value < 0:
            raise CriterionError(422, "response_out_of_range", f"{item_id}={value} must be >= 0")
        return
    scale = instrument["response_scale"]
    lo, hi = scale.get("min"), scale.get("max")
    if (lo is not None and value < lo) or (hi is not None and value > hi):
        raise CriterionError(422, "response_out_of_range", f"{item_id}={value} outside [{lo}, {hi}]")
    anchors = [a["value"] for a in scale.get("anchors", [])]
    if scale["type"] in ("likert", "pictorial_likert") and anchors and value not in anchors:
        raise CriterionError(422, "response_not_on_scale", f"{item_id}={value} is not a scale anchor")


def submit_administration(db: Session, adm: Administration, actor: str) -> tuple[list[Response], ScoringRun]:
    """Validate and store one administration, then score it.

    Respondent/subject rule: ``self`` -> pid is respondent and subject; ``partner`` -> pid is the
    rating participant and ``rated_pid`` the subject (must differ: self-rating is impossible);
    ``observer``/``administrator`` -> pid is the subject and the rater is ``administrator_id``.
    """
    if adm.site_mode == "open_public":
        raise CriterionError(
            422, "public_kiosk_context", "criterion records from a public kiosk context are rejected (CRT-04/CRT-05)"
        )
    role = adm.rater_role or "self"
    if role == "self" and adm.rated_pid not in (None, adm.pid):
        raise CriterionError(422, "rater_mismatch", "self ratings cannot target another participant")
    if role == "partner" and adm.rated_pid in (None, adm.pid):
        raise CriterionError(422, "self_rating_forbidden", "partner ratings need a different rated_pid")
    if role in ("observer", "administrator") and not adm.administrator_id:
        raise CriterionError(422, "rater_required", "observer/administrator ratings need administrator_id")
    subject = adm.rated_pid if role == "partner" else adm.pid
    assert subject is not None
    people = {adm.pid, subject}
    for person in sorted(people):
        participant = db.get(Participant, person)
        if participant is None:
            raise CriterionError(404, "unknown_pid", f"participant {person} not found")
        if participant.status != "active":
            raise CriterionError(409, "participant_withdrawn", "participant has withdrawn")
        if not consent_svc.has_consent(db, "criterion", pid=person):
            raise CriterionError(409, "criterion_consent_missing", f"no criterion consent recorded for {person}")
    version = get_version(db, adm.instrument_id, adm.instrument_version)
    assert_administrable(db, version, adm.language)
    definition = version.definition
    resolved: dict[str, tuple[dict[str, Any], str | None]] = {}
    for r in adm.responses:
        if r.item_id in resolved:
            raise CriterionError(422, "duplicate_item", f"{r.item_id} submitted twice")
        item, measure = _resolve_item(definition, r.item_id)
        resolved[r.item_id] = (item, measure)
        _validate_value(definition, r.item_id, r.raw_response)
        if measure == "moves" and r.raw_response is not None and r.expected_response is None:
            raise CriterionError(
                422,
                "performance_key_required",
                f"{r.item_id} needs expected_response = the problem's minimum moves (administration file)",
            )
        if r.expected_response is not None and not (item["attention_check"] or measure == "moves"):
            raise CriterionError(
                422, "unexpected_key", f"expected_response applies to attention checks and .moves only ({r.item_id})"
            )
    answered = {str(resolved[r.item_id][0]["item_id"]) for r in adm.responses if r.raw_response is not None}
    planned = _planned_missing(definition, adm.planned_missing_items, answered=answered)
    now = datetime.now(UTC)
    rows: list[Response] = []
    for r in adm.responses:
        item, measure = resolved[r.item_id]
        failed = _attention_failed(bool(item["attention_check"]), r.raw_response, r.expected_response)
        row = Response(
            response_id=uuid.uuid4(),
            pid=adm.pid,
            subject_pid=subject,
            study_id=adm.study_id,
            wave=adm.wave,
            occasion_id=adm.occasion_id,
            instrument_id=adm.instrument_id,
            instrument_version=version.version,
            language_version=adm.language.lower(),
            block=adm.block,
            item_id=r.item_id,
            raw_response=r.raw_response,
            response_latency_ms=r.response_latency_ms,
            administration_mode=adm.administration_mode,
            administrator_id=adm.administrator_id,
            rater_role=role,
            rated_pid=adm.rated_pid if role == "partner" else None,
            device_id=adm.device_id,
            audio_version=adm.audio_version,
            attention_check=bool(item["attention_check"]),
            attention_check_failed=failed,
            expected_response=r.expected_response,
            planned_missing_items=planned,
            source=adm.source,
            import_id=adm.import_id,
            recorded_at=now,
            submitted_by=actor,
        )
        db.add(row)
        rows.append(row)
    db.flush()
    run = rescore(
        db,
        pid=subject,
        study_id=adm.study_id,
        occasion_id=adm.occasion_id,
        instrument_id=adm.instrument_id,
        actor=actor,
    )
    return rows, run


def _planned_missing(definition: dict[str, Any], requested: list[str], *, answered: set[str]) -> list[str]:
    """Validated, sorted planned-missing item ids: registry items of this instrument that were not answered."""
    items = {i["item_id"] for i in definition["items"]}
    unknown = sorted(set(requested) - items)
    if unknown:
        raise CriterionError(422, "unknown_item", f"planned-missing items are not part of the instrument: {unknown}")
    clash = sorted(set(requested) & answered)
    if clash:
        raise CriterionError(422, "planned_missing_answered", f"items marked planned missing were answered: {clash}")
    return sorted(set(requested))


def _attention_failed(attention_check: bool, raw: float | None, expected: float | None) -> bool | None:
    """Attention-check outcome from the stored key; None when either value is missing."""
    if not attention_check or raw is None or expected is None:
        return None
    return raw != expected


def current_responses(db: Session, *, pid: str, study_id: str, occasion_id: str, instrument_id: str) -> list[Response]:
    rows = list(
        db.execute(
            select(Response)
            .where(
                Response.subject_pid == pid,
                Response.study_id == study_id,
                Response.occasion_id == occasion_id,
                Response.instrument_id == instrument_id,
            )
            .order_by(Response.created_at)
        ).scalars()
    )
    superseded = {r.supersedes_id for r in rows if r.supersedes_id}
    latest: dict[str, Response] = {}
    for r in rows:
        if r.response_id not in superseded:
            latest[r.item_id] = r
    return list(latest.values())


def attention_failed_for(db: Session, *, pid: str, study_id: str, occasion_id: str) -> bool:
    rows = db.execute(
        select(Response).where(
            Response.subject_pid == pid,
            Response.study_id == study_id,
            Response.occasion_id == occasion_id,
            Response.attention_check.is_(True),
            Response.rater_role == "self",
        )
    ).scalars()
    superseded: set[uuid.UUID] = set()
    all_rows = list(rows)
    for r in all_rows:
        if r.supersedes_id:
            superseded.add(r.supersedes_id)
    return any(r.attention_check_failed for r in all_rows if r.response_id not in superseded)


def rescore(
    db: Session,
    *,
    pid: str,
    study_id: str,
    occasion_id: str,
    instrument_id: str,
    actor: str,
    planned_missing_items: frozenset[str] = frozenset(),
    job_id: uuid.UUID | None = None,
) -> ScoringRun:
    """Score the current responses for one (pid, occasion, instrument) with the versioned key.

    Planned-missing items are the union of ``planned_missing_items`` and those stored with the responses
    (administration design), so corrections and key updates never turn a planned gap into non-response.
    """
    responses = current_responses(db, pid=pid, study_id=study_id, occasion_id=occasion_id, instrument_id=instrument_id)
    if not responses:
        raise CriterionError(404, "no_responses", "nothing to score")
    answered = {r.item_id.split(".", 1)[0] for r in responses if r.raw_response is not None}
    stored = {i for r in responses for i in (r.planned_missing_items or [])}
    planned_missing_items = frozenset((set(planned_missing_items) | stored) - answered)
    version = get_version(db, instrument_id, responses[0].instrument_version)
    definition = version.definition
    inputs = [
        ResponseInput(
            item_id=r.item_id,
            raw_response=r.raw_response,
            response_id=str(r.response_id),
            block=r.block,
            expected_response=r.expected_response,
            administration_mode=r.administration_mode,
        )
        for r in responses
    ]
    careless = attention_failed_for(db, pid=pid, study_id=study_id, occasion_id=occasion_id)
    scales, algorithm = score_instrument(
        definition, inputs, attention_failed=careless, planned_missing_items=planned_missing_items
    )
    run = ScoringRun(
        kind="criterion",
        algorithm_version=algorithm,
        engine=algorithm,
        dictionary_versions={"key_version": version.key_version},
        dictionary_sha256={"instrument_sha256": version.sha256},
        inputs={
            "pid_hash": sha256_hex(pid)[:16],
            "study_id": study_id,
            "occasion_id": occasion_id,
            "instrument_id": instrument_id,
            "instrument_version": version.version,
        },
        status="succeeded",
        job_id=job_id,
        finished_at=datetime.now(UTC),
    )
    db.add(run)
    db.flush()
    db.execute(
        update(ScaleScore)
        .where(
            ScaleScore.pid == pid,
            ScaleScore.study_id == study_id,
            ScaleScore.occasion_id == occasion_id,
            ScaleScore.instrument_id == instrument_id,
            ScaleScore.is_current.is_(True),
        )
        .values(is_current=False)
    )
    by_item = {r.item_id: r for r in responses}
    wave = responses[0].wave
    for sc in scales:
        db.add(
            ScaleScore(
                pid=pid,
                study_id=study_id,
                wave=wave,
                occasion_id=occasion_id,
                instrument_id=instrument_id,
                instrument_version=version.version,
                scale_id=sc.scale_id,
                score=sc.score,
                raw_mean=sc.raw_mean,
                score_type=sc.score_type,
                n_answered=sc.n_answered,
                n_items=sc.n_items,
                key_version=version.key_version,
                prorated=sc.prorated,
                planned_missing=sc.planned_missing,
                na_reason=sc.na_reason,
                flags=sc.flags,
                scoring_run_id=run.run_id,
                is_current=True,
                algorithm_version=algorithm,
                dictionary_version=version.key_version,
                source_refs={"response_ids": [str(by_item[i].response_id) for i in sc.source_item_ids if i in by_item]},
                job_id=job_id,
            )
        )
    run.manifest = {
        "scales": len(scales),
        "responses": len(responses),
        "attention_failed": careless,
        "planned_missing_items": sorted(planned_missing_items),
        "actor": actor,
    }
    db.flush()
    return run


def correct_response(
    db: Session, response_id: uuid.UUID, *, raw_response: float | None, reason: str, actor: str
) -> tuple[Response, ScoringRun]:
    original = db.get(Response, response_id)
    if original is None:
        raise CriterionError(404, "unknown_response", "response not found")
    if db.execute(select(Response.response_id).where(Response.supersedes_id == response_id)).first():
        raise CriterionError(409, "already_superseded", "correct the latest version of this response")
    version = get_version(db, original.instrument_id, original.instrument_version)
    _validate_value(version.definition, original.item_id, raw_response)
    failed = _attention_failed(original.attention_check, raw_response, original.expected_response)
    row = Response(
        response_id=uuid.uuid4(),
        pid=original.pid,
        subject_pid=original.subject_pid,
        study_id=original.study_id,
        wave=original.wave,
        occasion_id=original.occasion_id,
        instrument_id=original.instrument_id,
        instrument_version=original.instrument_version,
        language_version=original.language_version,
        block=original.block,
        item_id=original.item_id,
        raw_response=raw_response,
        response_latency_ms=original.response_latency_ms,
        administration_mode=original.administration_mode,
        administrator_id=original.administrator_id,
        rater_role=original.rater_role,
        rated_pid=original.rated_pid,
        device_id=original.device_id,
        audio_version=original.audio_version,
        attention_check=original.attention_check,
        attention_check_failed=failed,
        expected_response=original.expected_response,
        planned_missing_items=list(original.planned_missing_items or []),
        source=original.source,
        import_id=original.import_id,
        recorded_at=datetime.now(UTC),
        supersedes_id=original.response_id,
        correction_reason=reason,
        submitted_by=actor,
    )
    db.add(row)
    db.flush()
    run = rescore(
        db,
        pid=original.subject_pid,
        study_id=original.study_id,
        occasion_id=original.occasion_id,
        instrument_id=original.instrument_id,
        actor=actor,
    )
    if original.attention_check:
        # The careless flag covers every scale of the occasion: rescore the occasion's other instruments.
        others = db.execute(
            select(Response.instrument_id)
            .where(
                Response.subject_pid == original.subject_pid,
                Response.study_id == original.study_id,
                Response.occasion_id == original.occasion_id,
                Response.instrument_id != original.instrument_id,
            )
            .distinct()
        ).scalars()
        for instrument_id in sorted(others):
            rescore(
                db,
                pid=original.subject_pid,
                study_id=original.study_id,
                occasion_id=original.occasion_id,
                instrument_id=instrument_id,
                actor=actor,
            )
    return row, run


def response_history(db: Session, response_id: uuid.UUID) -> list[Response]:
    """The full supersession chain containing ``response_id`` (oldest first)."""
    row = db.get(Response, response_id)
    if row is None:
        raise CriterionError(404, "unknown_response", "response not found")
    chain = [row]
    cur = row
    while cur.supersedes_id is not None:
        prev = db.get(Response, cur.supersedes_id)
        if prev is None:
            break
        chain.insert(0, prev)
        cur = prev
    cur = row
    while True:
        nxt = db.execute(select(Response).where(Response.supersedes_id == cur.response_id)).scalar_one_or_none()
        if nxt is None:
            break
        chain.append(nxt)
        cur = nxt
    return chain


# ---------------------------------------------------------------- imports


EXPECTED_PREFIX = "expected:"


def _read_table(filename: str, data: bytes) -> tuple[str, list[dict[str, Any]]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = data.decode("utf-8-sig")
        return "csv", [dict(r) for r in csv.DictReader(io.StringIO(text))]
    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook  # type: ignore[import-untyped]  # noqa: PLC0415

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        rows = [dict(zip(header, r, strict=False)) for r in it if any(v is not None and v != "" for v in r)]
        return "xlsx", rows
    raise CriterionError(415, "unsupported_file", "only .csv and .xlsx files are supported")


def _to_number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() in ("", "NA", "N/A", "na", "-")):
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a response")
    return float(value)


def stage_import(
    db: Session,
    *,
    filename: str,
    data: bytes,
    study_id: str,
    instrument_id: str,
    instrument_version: str | None,
    language: str,
    occasion_id: str,
    wave: int,
    block: str,
    administration_mode: str,
    column_mapping: dict[str, str],
    actor: str,
) -> ImportBatch:
    """Parse, map and validate a CSV/XLSX file; ambiguous rows go to the manual-review queue."""
    fmt, records = _read_table(filename, data)
    version = get_version(db, instrument_id, instrument_version)
    definition = version.definition
    pid_col = column_mapping.get("pid")
    if not pid_col:
        raise CriterionError(422, "mapping_missing_pid", "column_mapping must map 'pid'")
    item_map = {k.split(":", 1)[1]: v for k, v in column_mapping.items() if k.startswith("item:")}
    # ``expected:<item_id>`` columns carry administration-file keys (attention answers, ToL min moves).
    expected_map = {k.split(":", 1)[1]: v for k, v in column_mapping.items() if k.startswith("expected:")}
    unknown: list[str] = []
    for item_id in item_map:
        try:
            _resolve_item(definition, item_id)
        except CriterionError:
            unknown.append(item_id)
    unknown += [i for i in expected_map if i not in item_map]
    if unknown:
        raise CriterionError(422, "mapping_unknown_items", "mapping references unknown items", items=sorted(unknown))
    if not item_map:
        raise CriterionError(422, "mapping_no_items", "map at least one item column (item:<item_id>)")
    headers = set(records[0].keys()) if records else set()
    missing_cols = sorted({pid_col, *item_map.values(), *expected_map.values()} - headers)
    if records and missing_cols:
        raise CriterionError(
            422, "mapping_missing_columns", "mapped columns are absent from the file", columns=missing_cols
        )
    batch = ImportBatch(
        study_id=study_id,
        instrument_id=instrument_id,
        instrument_version=version.version,
        language=language,
        occasion_id=occasion_id,
        wave=wave,
        block=block,
        administration_mode=administration_mode,
        filename=filename,
        file_format=fmt,
        file_sha256=sha256_hex(data),
        column_mapping=column_mapping,
        status="staged",
        created_by=actor,
    )
    db.add(batch)
    db.flush()
    seen_pids: dict[str, int] = {}
    counts = {"ready": 0, "needs_review": 0}
    for n, rec in enumerate(records, start=1):
        issues: list[dict[str, Any]] = []
        raw_pid = str(rec.get(pid_col) or "").strip()
        pid = normalise_code(raw_pid) if raw_pid else ""
        if not pid or not is_valid_code(pid):
            issues.append({"code": "invalid_pid", "column": pid_col, "value": raw_pid})
            pid_value: str | None = None
        else:
            pid_value = pid
            participant = db.get(Participant, pid)
            if participant is None:
                issues.append({"code": "unknown_pid", "column": pid_col})
            elif participant.status != "active":
                issues.append({"code": "participant_withdrawn", "column": pid_col})
            if pid in seen_pids:
                issues.append({"code": "duplicate_pid_in_file", "first_row": seen_pids[pid]})
            seen_pids.setdefault(pid, n)
            if participant is not None and current_responses(
                db, pid=pid, study_id=study_id, occasion_id=occasion_id, instrument_id=instrument_id
            ):
                issues.append({"code": "existing_responses", "detail": "pid already has responses for occasion"})
        values: dict[str, float | None] = {}
        for item_id, col in item_map.items():
            try:
                v = _to_number(rec.get(col))
                _validate_value(definition, item_id, v)
                values[item_id] = v
            except (ValueError, TypeError):
                issues.append({"code": "not_numeric", "column": col, "value": str(rec.get(col))})
            except CriterionError as exc:
                issues.append({"code": exc.code, "column": col, "value": str(rec.get(col))})
        for item_id, col in expected_map.items():
            try:
                values[f"{EXPECTED_PREFIX}{item_id}"] = _to_number(rec.get(col))
            except (ValueError, TypeError):
                issues.append({"code": "not_numeric", "column": col, "value": str(rec.get(col))})
        status = "needs_review" if issues else "ready"
        counts[status] += 1
        db.add(
            ImportRow(
                import_id=batch.import_id, row_number=n, pid=pid_value, values=values, status=status, issues=issues
            )
        )
    batch.counts = {**counts, "rows": len(records)}
    db.flush()
    return batch


def resolve_import_row(
    db: Session,
    batch: ImportBatch,
    row: ImportRow,
    *,
    action: str,
    actor: str,
    note: str | None,
    pid: str | None,
    values: dict[str, float | None] | None,
) -> ImportRow:
    if batch.status != "staged":
        raise CriterionError(409, "import_not_staged", "import is no longer staged")
    if action == "reject":
        row.status = "rejected"
    elif action == "approve":
        if pid is not None:
            code = normalise_code(pid)
            if not is_valid_code(code) or db.get(Participant, code) is None:
                raise CriterionError(422, "invalid_pid", "replacement pid is invalid or unknown")
            row.pid = code
        if values is not None:
            version = get_version(db, batch.instrument_id, batch.instrument_version)
            for item_id, v in values.items():
                if not item_id.startswith(EXPECTED_PREFIX):
                    _resolve_item(version.definition, item_id)
                    _validate_value(version.definition, item_id, v)
            row.values = {**row.values, **values}
        if row.pid is None:
            raise CriterionError(422, "pid_required", "an approved row needs a valid pid")
        row.status = "approved"
    else:
        raise CriterionError(422, "invalid_action", "action must be approve or reject")
    row.resolved_by = actor
    row.resolved_at = datetime.now(UTC)
    row.resolution_note = note
    db.flush()
    return row


def commit_import(db: Session, batch: ImportBatch, *, actor: str, device_id: str) -> dict[str, int]:
    if batch.status != "staged":
        raise CriterionError(409, "import_not_staged", "import is no longer staged")
    pending = (
        db.execute(select(ImportRow).where(ImportRow.import_id == batch.import_id, ImportRow.status == "needs_review"))
        .scalars()
        .first()
    )
    if pending is not None:
        raise CriterionError(409, "review_pending", "resolve every needs_review row before committing")
    rows: Iterable[ImportRow] = db.execute(
        select(ImportRow)
        .where(ImportRow.import_id == batch.import_id, ImportRow.status.in_(("ready", "approved")))
        .order_by(ImportRow.row_number)
    ).scalars()
    committed = 0
    source = "import_xlsx" if batch.file_format == "xlsx" else "import_csv"
    for row in rows:
        assert row.pid is not None
        adm = Administration(
            pid=row.pid,
            study_id=batch.study_id,
            wave=batch.wave,
            occasion_id=batch.occasion_id,
            instrument_id=batch.instrument_id,
            instrument_version=batch.instrument_version,
            language=batch.language,
            block=batch.block,
            administration_mode=batch.administration_mode,
            administrator_id=None,
            device_id=device_id,
            audio_version=None,
            site_mode="supervised",
            rater_role="self",
            rated_pid=None,
            source=source,
            planned_missing_items=[],
            responses=[
                ItemResponse(item_id=k, raw_response=v, expected_response=row.values.get(f"{EXPECTED_PREFIX}{k}"))
                for k, v in sorted(row.values.items())
                if not k.startswith(EXPECTED_PREFIX)
            ],
            import_id=batch.import_id,
        )
        submit_administration(db, adm, actor)
        row.status = "committed"
        committed += 1
    batch.status = "committed"
    batch.committed_at = datetime.now(UTC)
    batch.committed_by = actor
    batch.counts = {**batch.counts, "committed": committed}
    db.flush()
    return {"committed": committed}
